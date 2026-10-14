#!/usr/bin/env python3
"""Read-only, provenance-first extraction of exported Braintime experiment sheets.

This creates review observations, never asserts animal identity or experiment completion.
Original XLSX archives remain authoritative for formatting, comments and attachments.
"""
import argparse
import collections
import datetime as dt
import gzip
import hashlib
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

NS = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
ASSAYS = [('LumiCycle', r'lumi\s*cyc'), ('IMG', r'\bimg\b|imag'), ('qPCR', r'\bqpcr\b'),
          ('WB', r'\bwb\b|western'), ('IF/IHC', r'\bihc\b|\bif\b|immunohisto'),
          ('CSF', r'\bcsf\b|cerebrospinal'), ('tissue_storage', r'[-−]\s*(?:70|80)\s*[º°]?\s*c|freez|fridge')]


def js(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool, list, dict)):
        return value
    return str(value)


def text(value):
    return '' if value is None else str(js(value)).strip()


def meaningful(value):
    return not isinstance(value, bool) and value is not None and text(value) != ''


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), default=js).encode()).hexdigest()


def date_value(value):
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.year >= 1900 else None
    if isinstance(value, dt.date):
        return value.isoformat()
    raw = text(value)
    if re.search(r'(?:19|20)\d{2}[-/.]\d{1,2}[-/.]\d{1,2}\s*[-–~]\s*\d{1,2}', raw):
        return None
    matches = re.findall(r'(?<!\d)((?:19|20)\d{2})[-/.](\d{1,2})[-/.](\d{1,2})(?!\d)', raw)
    if len(set(matches)) != 1:
        return None
    try:
        return dt.date(*map(int, matches[0])).isoformat()
    except ValueError:
        return None


def assay_labels(value):
    return [name for name, pattern in ASSAYS if re.search(pattern, text(value), re.I)]


def identifiers(value):
    """Preserve source labels; do not turn tissue IDs, cages or tags into mouse SIDs."""
    if not meaningful(value):
        return []
    raw = text(value)
    tags = re.findall(r"(?<![A-Za-z0-9])((?:AD|WT|B6|PER|Vg|Gt|32|96|[PEBFVA'′’()])+(?:[- ]+[MF]\s*-?\s*\d+|[MF]\d+))(?!\d)", raw)
    if tags:
        return list(dict.fromkeys(tag.strip() for tag in tags))
    return list(dict.fromkeys(s.strip() for s in re.split(r'[,;\n]+', raw) if s.strip()))


def normalized(value):
    return re.sub(r'\s+', ' ', text(value)).lower()


def sheet_xml(archive, index):
    root = ET.fromstring(archive.read(f'xl/worksheets/sheet{index + 1}.xml'))
    merges = [m.attrib['ref'] for m in root.findall('x:mergeCells/x:mergeCell', NS)]
    formulas = {c.attrib['r']: c.find('x:f', NS).text or '' for c in root.findall('.//x:c', NS) if c.find('x:f', NS) is not None}
    links = {}
    relpath = f'xl/worksheets/_rels/sheet{index + 1}.xml.rels'
    rels = {}
    if relpath in archive.namelist():
        rels = {r.attrib['Id']: r.attrib.get('Target', '') for r in ET.fromstring(archive.read(relpath))}
    for h in root.findall('x:hyperlinks/x:hyperlink', NS):
        links[h.attrib['ref']] = rels.get(h.attrib.get(REL), h.attrib.get('location', ''))
    return merges, formulas, links


def merged_values(rows, merges):
    resolved = [list(r) for r in rows]
    merged = {}
    for ref in merges:
        minc, minr, maxc, maxr = range_boundaries(ref)
        if minr > len(rows):
            continue
        value = rows[minr - 1][minc - 1] if minc <= len(rows[minr - 1]) else None
        if value is None:
            continue
        for r in range(minr, min(maxr, len(rows)) + 1):
            if len(resolved[r - 1]) < maxc:
                resolved[r - 1].extend([None] * (maxc - len(resolved[r - 1])))
            for c in range(minc, maxc + 1):
                if r == minr and c == minc:
                    continue
                if resolved[r - 1][c - 1] is None:
                    resolved[r - 1][c - 1] = value
                    merged[f'{get_column_letter(c)}{r}'] = ref
    return resolved, merged


def at(row, index):
    return row[index] if index is not None and 0 <= index < len(row) else None


def field_index(headers, *patterns):
    for pattern in patterns:
        for i, value in enumerate(headers):
            if re.search(pattern, normalized(value)):
                return i
    return None


def row_fields(headers, row):
    return {f'{get_column_letter(i + 1)} · {text(at(headers, i)) or "Column " + get_column_letter(i + 1)}': js(v)
            for i, v in enumerate(row) if v is not None and text(v) != ''}


def common_obs(kind, assays, raw_id, date, tissue='', storage='', raw=None, **extra):
    return dict(kind=kind, assays=assays, raw_identifiers=identifiers(raw_id), date=date_value(date),
                date_raw=text(date), tissue=text(tissue), storage=text(storage), raw_values=raw or {},
                animal_link_status='unresolved', review_status='unreviewed', **extra)


def parse_sheet(title, rows, resolved, workbook_title=''):
    """Yield (range, observation). Data selection is conservative and layout-specific."""
    low = title.strip().lower()
    if low == 'fpb img data':
        for ri, row in enumerate(resolved[1:], 2):
            if not meaningful(at(row, 4)):
                continue
            yield f'A{ri}:G{ri}', common_obs('imaging_sample', ['IMG'], at(row, 4), at(row, 0), at(row, 5),
                raw=row_fields(rows[0], row), identifier_kind='sample_label', channel=text(at(row, 6)))
        return
    header_i = next((i for i, r in enumerate(rows[:12]) if any('date (prep)' in normalized(v) for v in r)), None)
    # Instrument rows consist of one preparation block and per-channel vertical fields.
    if header_i is not None and ('lumicycle' in low or 'imaging' in low or 'kronos' in low):
        headers = rows[header_i]
        label_col = next((j for r in rows[header_i + 1:header_i + 10] for j, v in enumerate(r)
                          if normalized(v) in ('animal id', 'gender / #')), None)
        if label_col is not None:
            channels = [j for j, h in enumerate(headers) if re.search(r'^\d+\s*\(\d+[a-z]\)|^imaging\s*\d|^kronos', normalized(h))]
            if not channels:
                channels = list(range(label_col + 1, min(len(headers), label_col + 5)))
            starts = [i for i in range(header_i + 1, len(rows)) if normalized(at(rows[i], label_col)) in ('animal id', 'gender / #')]
            starts.append(len(rows))
            for start, stop in zip(starts, starts[1:]):
                main = resolved[start]
                dateidx = field_index(headers, r'^date \(prep\)')
                for col in channels:
                    channel = {text(at(rows[r], label_col)): js(at(resolved[r], col)) for r in range(start, min(stop, start + 15))
                               if meaningful(at(rows[r], label_col)) and at(resolved[r], col) is not None}
                    animal = at(main, col)
                    tissue = next((v for k, v in channel.items() if normalized(k) == 'tissue'), '')
                    if not meaningful(animal) and not meaningful(tissue):
                        continue
                    raw = {'channel': text(at(headers, col)), 'sample': channel,
                           'preparation': {k: v for k, v in row_fields(headers, main).items()
                                           if k.split(' · ')[0] not in {get_column_letter(j + 1) for j in channels}}}
                    assay = 'LumiCycle' if 'lumicycle' in low else 'IMG' if 'imaging' in low else 'Other'
                    yield f'{get_column_letter(col + 1)}{start + 1}:{get_column_letter(col + 1)}{min(stop, start + 15)}', common_obs(
                        'instrument_channel', [assay], animal, at(main, dateidx), tissue, raw=raw,
                        channel=text(at(headers, col)), operator=text(at(main, field_index(headers, r'^user$'))),
                        project=text(at(main, field_index(headers, r'^project$'))))
            return
    # OD CSF columns are individual animals, each with labelled metadata rows.
    if low == 'od csf':
        idrows = [i for i, r in enumerate(rows) if normalized(at(r, 0)) == 'mouse id']
        for idrow in idrows:
            start = max(0, idrow - 5)
            stop = min(len(rows), next((i for i in range(idrow + 1, len(rows)) if normalized(at(rows[i], 0)) == 'zt'), len(rows)))
            fields = [(i, normalized(at(rows[i], 0))) for i in range(start, stop) if meaningful(at(rows[i], 0))]
            dateidx = next((i for i, label in fields if label == 'sampling date'), None)
            for col in range(1, len(rows[idrow])):
                raw_id = at(rows[idrow], col)
                if not meaningful(raw_id):
                    continue
                raw = {text(at(rows[i], 0)): js(at(rows[i], col)) for i, label in fields if at(rows[i], col) is not None}
                yield f'{get_column_letter(col + 1)}{start + 1}:{get_column_letter(col + 1)}{stop}', common_obs(
                    'csf_sample', ['CSF'], raw_id, at(rows[dateidx], col) if dateidx is not None else None,
                    'CSF', raw=raw, phase=text(raw.get('Actual ZT (CSF ZT)', raw.get('ZT', ''))))
        return
    # Tissue sampling and qPCR usage tables use headers; actual merged cells are resolved.
    if low in ('sampling', 'sampling (qpcr/wb/ihc)', 'qpcr') or 'tissue sampling' in low:
        header_i = next((i for i, r in enumerate(rows[:12]) if any(re.search(r'^date|^qpcr\s*date', normalized(v)) for v in r)
                        and any(re.search(r'tissue|sample|animals', normalized(v)) for v in r)), None)
        if header_i is None:
            return
        headers = list(rows[header_i])
        # qPCR has a second header row for sample handling and assay fields.
        second = rows[header_i + 1] if header_i + 1 < len(rows) else []
        if low == 'qpcr' and any('purification date' in normalized(v) for v in second):
            headers = [text(at(headers, i)) + (' / ' + text(v) if meaningful(v) else '') for i, v in enumerate(second)]
            header_end = header_i + 2
        else:
            header_end = header_i + 1
        dateidx = field_index(headers, r'^date', r'qpcr\s*date')
        ididx = field_index(headers, r'^name tag$', r'^animals')
        animalidx = field_index(headers, r'^animals')
        tissueidx = field_index(headers, r'^tissues$', r'^tissue sample$', r'^sample$')
        storageidx = field_index(headers, r'^storage$')
        assaysidx = field_index(headers, r'qpcr\s*/\s*wb')
        for ri in range(header_end, len(rows)):
            row = resolved[ri]
            # Exclude prepared checkbox/number rows; require scientific content.
            if not any(meaningful(at(rows[ri], j)) for j in [ididx, tissueidx, dateidx] if j is not None):
                continue
            if normalized(at(row, dateidx)).startswith('date'):
                continue
            raw = row_fields(headers, row)
            storage = at(row, storageidx)
            if not storage:
                storage = '\n'.join(text(v) for v in row if re.search(r'[-−]\s*(?:70|80)\s*[º°]?\s*c|freez', text(v), re.I))
            assays = assay_labels(at(row, assaysidx)) if assaysidx is not None else (['qPCR'] if low == 'qpcr' else assay_labels(' '.join(map(text, row))))
            assays += [a for a in assay_labels(storage) if a not in assays]
            if not assays:
                assays = ['Other']
            yield f'A{ri + 1}:{get_column_letter(max(1, len(row)))}{ri + 1}', common_obs(
                'tissue_sampling' if low != 'qpcr' else 'assay_run', assays,
                '\n'.join(dict.fromkeys(text(at(row, i)) for i in (ididx, animalidx) if i is not None and meaningful(at(row, i)))), at(row, dateidx),
                at(row, tissueidx), storage, raw, operator=text(at(row, field_index(headers, r'^sampler$', r'^user$'))),
                project=text(at(row, field_index(headers, r'^project$'))), phase=text(at(row, field_index(headers, r'\(zt\)'))),
                time=text(at(row, field_index(headers, r'^sampling hour$'))))
        return
    # WB and IHC name cells can be tissue/sample labels, never automatically mouse IDs.
    if low in ('wb', 'ihc'):
        top = next((i for i, r in enumerate(rows[:8]) if normalized(at(r, 0)) in ('running time', 'performing date')), None)
        if top is None:
            return
        sub = rows[top + 1]
        headers = [text(at(rows[top], i)) + ' / ' + text(v) for i, v in enumerate(sub)]
        nameidx = field_index(sub, r'^name$')
        subidx = field_index(sub, r'^subregion$')
        for ri in range(top + 2, len(rows)):
            row = resolved[ri]
            if not any(meaningful(at(rows[ri], j)) for j in (0, 1, nameidx, subidx) if j is not None):
                continue
            if normalized(at(row, 0)) in ('running time', 'performing date', 'date') or normalized(at(row, nameidx)) == 'name':
                continue
            yield f'A{ri + 1}:{get_column_letter(max(1, len(row)))}{ri + 1}', common_obs(
                'assay_run', ['WB' if low == 'wb' else 'IF/IHC'], at(row, nameidx), at(row, 0),
                at(row, subidx), raw=row_fields(headers, row), sampling_date=date_value(at(row, 1)),
                identifier_kind='sample_label', operator=text(at(row, field_index(rows[top], r'^user$'))),
                project=text(at(row, field_index(rows[top], r'^project$'))))
        return
    # Preserve other relevant matrix and supplementary evidence as review rows.
    relevant = re.search(r'sampling|sample|fridge|fpb img|lumicycle|kronos|wb|ihc', low) or 'WB-CPs' in workbook_title
    if relevant and not re.search(r'illuminance', low):
        for ri, row in enumerate(rows):
            vals = [v for v in row if meaningful(v)]
            if len(vals) < 2 or ri < 1:
                continue
            labels = assay_labels(title + ' ' + workbook_title)
            if not labels:
                labels = assay_labels(' '.join(map(text, vals))) or ['Other']
            yield f'A{ri + 1}:{get_column_letter(len(row))}{ri + 1}', common_obs(
                'source_matrix_review', labels, '', None, raw=row_fields([], row),
                parsing_note='Matrix or supplementary evidence: inspect the original source before creating animal associations.')


def build(source_dir, output):
    source_dir = Path(source_dir)
    inventories = [json.loads((source_dir / name).read_text())['results'] for name in ('experiment_inventory.json', 'supplementary_inventory.json')]
    files = {f['id']: f for inventory in inventories for f in inventory}
    observations = {}
    sources, sheet_reports = [], []
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    rawout = gzip.open(output / 'experiment_raw_rows.jsonl.gz', 'wt', encoding='utf-8')
    for file_id, file in files.items():
        path = source_dir / (file_id + '.xlsx')
        meta_path = source_dir / 'metadata' / (file_id + '.json')
        if not meta_path.exists():
            current_meta = source_dir / 'current_metadata.json'
            if current_meta.exists() and json.loads(current_meta.read_text()).get('spreadsheetId') == file_id:
                meta_path = current_meta
        metadata = json.loads(meta_path.read_text())
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        source = dict(file_id=file_id, title=file['title'], url=file['url'], archive_path=str(path.resolve()),
                      archive_sha256=sha, archive_size=path.stat().st_size, snapshot_date=date_value(file['title']),
                      snapshot_date_basis='title_date_not_verified_capture' if date_value(file['title']) else 'capture_date',
                      file_created_at=file.get('created_at'), file_modified_at=file.get('updated_at'),
                      exported_at=dt.datetime.now(dt.timezone.utc).isoformat())
        if source['snapshot_date'] is None:
            source['snapshot_date'] = dt.datetime.now(dt.timezone.utc).date().isoformat()
        sources.append(source)
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        archive = ZipFile(path)
        occurrences = collections.Counter()
        for si, sheet in enumerate(book):
            prop = metadata['sheets'][si]['properties']
            title, sid = prop['title'], prop['sheetId']
            rows = [list(r) for r in sheet.iter_rows(values_only=True)]
            merges, formulas, hyperlinks = sheet_xml(archive, si)
            formulas_by_row, links_by_row = collections.defaultdict(dict), collections.defaultdict(dict)
            for address, formula in formulas.items():
                formulas_by_row[int(re.sub(r'^[A-Z]+', '', address))][address] = formula
            for address, link in hyperlinks.items():
                match = re.search(r'\d+', address)
                if match:
                    links_by_row[int(match.group())][address] = link
            resolved, merged = merged_values(rows, merges)
            raw_count = 0
            for ri, row in enumerate(rows, 1):
                values = {get_column_letter(c + 1): js(v) for c, v in enumerate(row) if v is not None and text(v) != ''}
                if not values:
                    continue
                raw_count += 1
                rawout.write(json.dumps(dict(workbook_id=file_id, sheet_id=sid, sheet_name=title, export_title=sheet.title,
                                             row=ri, values=values,
                                             formulas=formulas_by_row[ri], hyperlinks=links_by_row[ri]),
                                        ensure_ascii=False, default=js, separators=(',', ':')) + '\n')
            parsed_count = 0
            for cell_range, obs in parse_sheet(title, rows, resolved, file['title']):
                parsed_count += 1
                # Keep repeat observations within one sheet distinct; coalesce exact copies across snapshots.
                fingerprint = digest(obs)
                occurrence_key = (title.strip(), fingerprint)
                occurrences[occurrence_key] += 1
                key = digest([title.strip(), fingerprint, occurrences[occurrence_key]])
                source_url = f'https://docs.google.com/spreadsheets/d/{file_id}/edit#gid={sid}&range={cell_range}'
                provenance = dict(workbook_id=file_id, workbook_title=file['title'], sheet_id=sid, sheet_name=title,
                                  range=cell_range, source_url=source_url, snapshot_date=source['snapshot_date'], archive_sha256=sha)
                if key in observations:
                    observations[key]['provenance'].append(provenance)
                    continue
                observations[key] = dict(source_id='experiment-source-' + key, content_hash=fingerprint,
                                         **provenance, **obs, provenance=[provenance])
            sheet_reports.append(dict(workbook_id=file_id, sheet_id=sid, sheet_name=title, export_title=sheet.title,
                                      grid=prop.get('gridProperties'), hidden=bool(prop.get('hidden')), populated_rows=raw_count,
                                      observation_occurrences=parsed_count, merged_ranges=merges,
                                      parser='structured' if parsed_count and not all(o['kind']=='source_matrix_review' for _,o in parse_sheet(title, rows, resolved, file['title'])) else 'review_matrix' if parsed_count else 'archive_only'))
        book.close()
        archive.close()
        print(file['title'], 'complete', flush=True)
    rawout.close()
    values = list(observations.values())
    result = dict(schema_version='1.0', imported_at=dt.datetime.now(dt.timezone.utc).isoformat(), sources=sources,
                  observations=values, coverage=dict(workbooks=len(sources), sheets=len(sheet_reports),
                  unique_observations=len(values), observation_occurrences=sum(s['observation_occurrences'] for s in sheet_reports),
                  sheets_by_parser=dict(collections.Counter(s['parser'] for s in sheet_reports)),
                  observations_by_kind=dict(collections.Counter(o['kind'] for o in values))),
                  caveats=['Source observations require review and do not assert experiment completion or animal identity.',
                           'Tags and sample labels are retained verbatim; repeated historical tags must not be linked automatically.',
                           'Only actual merged cells are resolved. Other missing dates remain unknown.',
                           'Recorded storage temperatures, including -70 C and -80 C, remain unchanged.',
                           'Exact semantic copies are coalesced across workbook snapshots, with every source retained; changed observations remain separate.',
                           'All exported tabs are preserved in raw rows and original XLSX, including hidden and archive-only tabs.',
                           'Exported workbook snapshots do not constitute complete Google Sheets edit-version history.',
                           'XLSX names may be sanitized; live metadata sheet IDs and titles are used by index.'])
    (output/'experiment_sources.json').write_text(json.dumps(result, ensure_ascii=False, default=js, separators=(',', ':')))
    (output/'sheet_coverage.json').write_text(json.dumps(sheet_reports, ensure_ascii=False, indent=2))
    (output/'coverage.json').write_text(json.dumps(result['coverage'], indent=2))
    print(json.dumps(result['coverage'], indent=2))


def add_animal_history(bundle_path, history_path):
    """Bring retained Animal Sheet usage evidence into the same review queue."""
    path = Path(bundle_path)
    bundle = json.loads(path.read_text())
    history = json.loads(Path(history_path).read_text())
    by_sheet = collections.defaultdict(lambda: collections.defaultdict(dict))
    for cell in history['cells']:
        match = re.fullmatch(r'([A-Z]+)(\d+)', cell[1])
        if match:
            by_sheet[cell[0]][int(match[2])][match[1]] = cell[3] if cell[3] is not None else cell[2]
    observations = {}
    occurrences = collections.Counter()
    source_ids = set()
    matched_cells = 0
    for si, sheet in enumerate(history['sheets']):
        source = history['sources'][sheet['source_index']]
        for rownum, cells in by_sheet[si].items():
            labels = assay_labels(' '.join(map(text, cells.values())))
            if not labels:
                continue
            is_usage = sheet['title'].strip().lower() == 'mouse usage' and rownum > 3
            # Generic cage notes are evidence, never proof of use for the current occupant.
            raw_id = cells.get('D', '') if is_usage else ''
            date = cells.get('A') if is_usage else None
            obs = common_obs('animal_usage' if is_usage else 'animal_sheet_review', labels, raw_id, date,
                             raw=cells, operator=text(cells.get('I')) if is_usage else '',
                             parsing_note='Animal Sheet usage record; verify identity and experiment completion.' if is_usage
                             else 'Animal Sheet context only; cage or tag mentions do not establish animal use.')
            fingerprint = digest(obs)
            occurrence_key = (source['file_id'], sheet['title'], fingerprint)
            occurrences[occurrence_key] += 1
            key = digest(['animal-sheet', sheet['title'], fingerprint, occurrences[occurrence_key]])
            last_col = max(cells, key=lambda x: (len(x), x))
            cell_range = f'A{rownum}:{last_col}{rownum}'
            ref = dict(workbook_id=source['file_id'], workbook_title=source['title'], sheet_id=sheet['sheet_id'],
                       sheet_name=sheet['title'], range=cell_range,
                       source_url=f'https://docs.google.com/spreadsheets/d/{source["file_id"]}/edit#gid={sheet["sheet_id"]}&range={cell_range}',
                       snapshot_date=source['snapshot_date'], archive_sha256=source['archive_sha256'])
            source_ids.add(source['file_id'])
            matched_cells += 1
            if key in observations:
                observations[key]['provenance'].append(ref)
            else:
                observations[key] = dict(source_id='experiment-source-'+key, content_hash=fingerprint,
                                         **ref, **obs, provenance=[ref])
    bundle['observations'] += list(observations.values())
    bundle['sources'] += [s for s in history['sources'] if s['file_id'] in source_ids]
    bundle['coverage']['animal_history_workbooks'] = len(source_ids)
    bundle['coverage']['animal_history_observations'] = len(observations)
    bundle['coverage']['animal_history_occurrences'] = matched_cells
    bundle['coverage']['unique_observations'] = len(bundle['observations'])
    path.write_text(json.dumps(bundle, ensure_ascii=False, separators=(',', ':')))
    print('Added Animal Sheet history:', len(observations), 'unique observations from', len(source_ids), 'workbooks')


def compact_bundle(bundle_path, output_path):
    """Preserve shared preparation metadata once and use catalog-indexed source refs."""
    bundle = json.loads(Path(bundle_path).read_text())
    sources = [{k: v for k, v in source.items() if k != 'archive_path'} for source in bundle['sources']]
    source_index = {source['file_id']: i for i, source in enumerate(sources)}
    contexts, observations = {}, []
    for original in bundle['observations']:
        obs = {k: v for k, v in original.items() if k not in ('provenance', 'source_url', 'workbook_title',
               'archive_sha256', 'content_hash', 'animal_link_status', 'review_status', 'parsing_note') and v not in ('', None, [])}
        obs['provenance_refs'] = [[source_index[ref['workbook_id']], ref['sheet_id'], ref['sheet_name'], ref['range']]
                                  for ref in original['provenance']]
        raw = dict(obs.get('raw_values', {}))
        if 'preparation' in raw:
            preparation = raw.pop('preparation')
            context_id = digest(preparation)
            contexts[context_id] = dict(context_id=context_id, values=preparation)
            obs['context_id'] = context_id
        obs['raw_values'] = raw
        observations.append(obs)
    coverage = dict(bundle['coverage'])
    coverage.update(total_source_workbooks=len(sources), unique_observations=len(observations),
                    preparation_contexts=len(contexts), review_status='unreviewed_source_observations',
                    full_history_is_edit_revisions=False,
                    observations_by_kind=dict(collections.Counter(o['kind'] for o in observations)))
    compact = dict(schema_version='1.1', imported_at=bundle['imported_at'], sources=sources, contexts=list(contexts.values()),
                   observations=observations, coverage=coverage, caveats=bundle['caveats'])
    payload = json.dumps(compact, ensure_ascii=False, separators=(',', ':'))
    Path(output_path).write_text(payload)
    with gzip.open(str(output_path)+'.gz', 'wt', encoding='utf-8') as out:
        out.write(payload)
    print('Compact bundle bytes:', len(payload.encode()), 'max observation bytes:',
          max(len(json.dumps(o, ensure_ascii=False, separators=(',', ':')).encode()) for o in observations))
    return compact


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source_dir')
    parser.add_argument('--output', required=True)
    parser.add_argument('--animal-history', help='Existing provenance-preserving Animal Sheet history_cells.json')
    args = parser.parse_args()
    build(args.source_dir, args.output)
    if args.animal_history:
        add_animal_history(Path(args.output)/'experiment_sources.json', args.animal_history)
    compact_bundle(Path(args.output)/'experiment_sources.json', Path(args.output)/'experiment_sources_app.json')
